"""Shared python-pptx helpers for the forecast-teller decks (epidemic dataviz palette, PingFang TC).

Usage: `from deck_lib import *` inside scripts/; call `new_deck()` first, then the helpers take a slide.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt, Emu

C = {"p50": "F6F9F6", "p100": "E8EEE7", "p300": "B4C9B1", "p500": "739A6D", "p600": "5D7F58", "p800": "374C34", "p900": "253423",
     "n0": "FFFFFF", "n100": "F2F3F1", "n200": "E4E7E4", "n300": "CACFC9", "n400": "A2ABA0", "n500": "7A8778", "n600": "5D675B", "n700": "444C43", "n800": "2C312B", "n900": "181B18",
     "slate": "587A9D", "mustard": "A8821F", "alert": "BE373C"}
FONT = "PingFang TC"
fmt0 = lambda v: f"{v:,.0f}"


def rgb(h): return RGBColor.from_string(h)


def new_deck():
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs, prs.slide_layouts[6]


def _ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rPr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, size=14, color="n800", bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, space_after=4, line=1.15, font=FONT):
    """paras: str | list of str | list of list-of-runs; a run is str or (str, {size,bold,color})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line
        runs = para if isinstance(para, list) else [para]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run(); run.text = t; f = run.font
            f.name = o.get("font", font); f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.color.rgb = rgb(C[o.get("color", color)]); _ea(run)
        if bullets:
            pPr = p._p.get_or_add_pPr(); pPr.set("marL", str(Emu(Inches(0.22)))); pPr.set("indent", str(-Emu(Inches(0.22))))
            bu = OxmlElement("a:buChar"); bu.set("char", "•"); pPr.append(bu)
    return tb


def rect(slide, x, y, w, h, fill="n100", line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb = rgb(C[fill])
    if line: s.line.color.rgb = rgb(C[line]); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False; return s


def circle_num(slide, x, y, n, d=0.36):
    c = rect(slide, x, y, d, d, fill="p500", shape=MSO_SHAPE.OVAL); tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(n); r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = rgb(C["n0"]); r.font.name = FONT; _ea(r)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def title(slide, t, sub=None, kicker=None):
    if kicker: text(slide, 0.6, 0.38, 8, 0.25, kicker, size=10, color="p600", bold=True)
    text(slide, 0.6, 0.6, 12.1, 0.6, t, size=26, color="p900", bold=True)
    if sub: text(slide, 0.6, 1.22, 12.1, 0.35, sub, size=12.5, color="n600")


def footer(slide, n, label):
    text(slide, 0.6, 7.05, 10, 0.25, label, size=8.5, color="n500")
    text(slide, 12.2, 7.05, 0.5, 0.25, str(n), size=9, color="n500", align=PP_ALIGN.RIGHT)


def table(slide, x, y, w, rows, col_w, header_fill="n100", font_size=10.5, best_row=None, row_h=0.3, num_from=None):
    """num_from: index of the first numeric (right-aligned) column; None = all columns left-aligned."""
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))); tbl = shp.table
    tblPr = tbl._tbl.tblPr; tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None: tblPr.remove(style)
    for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j); cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(C[header_fill if i == 0 else ("p50" if i == best_row else "n0")])
            tf = cell.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT if (num_from is not None and j >= num_from) else PP_ALIGN.LEFT
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            r = p.add_run(); r.text = str(val); r.font.size = Pt(font_size); r.font.name = FONT; _ea(r)
            r.font.bold = (i == 0) or (i == best_row and j == 0); r.font.color.rgb = rgb(C["n600" if i == 0 else "n800"])
            tcPr = cell._tc.get_or_add_tcPr()
            for k, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):  # must precede the fill in a:tcPr
                ln = OxmlElement(side); ln.set("w", "6350" if side == "a:lnB" else "0"); ln.set("cap", "flat"); ln.set("cmpd", "sng"); ln.set("algn", "ctr")
                if side == "a:lnB":
                    sf = OxmlElement("a:solidFill"); clr = OxmlElement("a:srgbClr"); clr.set("val", C["n200"]); sf.append(clr); ln.append(sf)
                else:
                    ln.append(OxmlElement("a:noFill"))
                tcPr.insert(k, ln)
    return tbl


def kpi_column(slide, x, y, w, items, step=0.78):
    """items: list of (label, value, sub)."""
    for i, (lab, val, sub) in enumerate(items):
        yy = y + i * step
        text(slide, x, yy, w, 0.25, lab, size=9.5, color="n500", bold=True)
        text(slide, x, yy + 0.24, w, 0.3, val, size=13, color="n900", bold=True)
        text(slide, x, yy + 0.53, w, 0.22, sub, size=9, color="n600")


__all__ = ["C", "FONT", "fmt0", "rgb", "new_deck", "text", "rect", "circle_num", "title", "footer", "table", "kpi_column",
           "Inches", "Pt", "PP_ALIGN", "MSO_ANCHOR", "MSO_SHAPE"]
